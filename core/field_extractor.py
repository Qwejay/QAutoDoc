#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
字段提取器
用于从Word文档中提取 {{字段名}} 格式的动态字段
"""

import re
from docx import Document


class FieldExtractor:
    """字段提取器类"""
    
    def __init__(self):
        # 正则表达式匹配 {{字段名}} 格式
        self.field_pattern = re.compile(r'\{\{(.*?)\}\}')
        
    def extract_fields(self, doc_path):
        """
        从Word文档中提取字段
        
        Args:
            doc_path: Word文档路径
            
        Returns:
            list: 字段名列表（按照文档中出现的顺序）
        """
        try:
            doc = Document(doc_path)
            fields = []  # 使用列表保持顺序
            seen_fields = set()  # 用于去重
            
            # 遍历文档中的所有段落
            for paragraph in doc.paragraphs:
                paragraph_fields = self._extract_from_text(paragraph.text)
                for field in paragraph_fields:
                    if field not in seen_fields:
                        fields.append(field)
                        seen_fields.add(field)
                
            # 遍历文档中的所有表格
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            cell_fields = self._extract_from_text(paragraph.text)
                            for field in cell_fields:
                                if field not in seen_fields:
                                    fields.append(field)
                                    seen_fields.add(field)
                            
            return fields
            
        except Exception as e:
            raise Exception(f"提取字段时出错：{str(e)}")
            
    def _extract_from_text(self, text):
        """
        从文本中提取字段
        
        Args:
            text: 输入文本
            
        Returns:
            list: 字段名列表（按照文本中出现的顺序）
        """
        matches = self.field_pattern.findall(text)
        # 保持字段在文本中出现的顺序，同时去重
        seen = set()
        fields = []
        for match in matches:
            field = match.strip()
            if field and field not in seen:
                fields.append(field)
                seen.add(field)
        return fields
    
    def extract_fields_with_positions(self, doc_path):
        """
        提取字段及其位置信息（用于高级替换功能）
        
        Args:
            doc_path: Word文档路径
            
        Returns:
            list: 包含字段位置信息的字典列表
        """
        try:
            doc = Document(doc_path)
            field_positions = []
            
            # 遍历段落
            for para_idx, paragraph in enumerate(doc.paragraphs):
                positions = self._find_field_positions_in_paragraph(paragraph.text, para_idx)
                field_positions.extend(positions)
                
            # 遍历表格
            for table_idx, table in enumerate(doc.tables):
                for row_idx, row in enumerate(table.rows):
                    for cell_idx, cell in enumerate(row.cells):
                        for para_idx, paragraph in enumerate(cell.paragraphs):
                            positions = self._find_field_positions_in_paragraph(
                                paragraph.text, 
                                para_idx,
                                table_idx=table_idx,
                                row_idx=row_idx,
                                cell_idx=cell_idx
                            )
                            field_positions.extend(positions)
                            
            return field_positions
            
        except Exception as e:
            raise Exception(f"提取字段位置时出错：{str(e)}")
            
    def _find_field_positions_in_paragraph(self, text, para_idx, table_idx=None, row_idx=None, cell_idx=None):
        """
        在段落文本中查找字段位置
        
        Args:
            text: 段落文本
            para_idx: 段落索引
            table_idx: 表格索引（可选）
            row_idx: 行索引（可选）
            cell_idx: 单元格索引（可选）
            
        Returns:
            list: 字段位置信息列表
        """
        positions = []
        matches = list(self.field_pattern.finditer(text))
        
        for match in matches:
            field_info = {
                'field_name': match.group(1).strip(),
                'start_pos': match.start(),
                'end_pos': match.end(),
                'full_match': match.group(0),
                'para_idx': para_idx,
                'table_idx': table_idx,
                'row_idx': row_idx,
                'cell_idx': cell_idx
            }
            positions.append(field_info)
            
        return positions