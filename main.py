#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
QAutoDoc - 智能自适应批量文档生成器 (Excel级数据矩阵强化版)
支持格式: DOCX / PPTX / XLSX
"""

import sys
import os
import re
import csv
import logging
import platform
import subprocess
from datetime import datetime
from pathlib import Path

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QVBoxLayout, QHBoxLayout, QWidget,
    QPushButton, QLabel, QLineEdit, QFileDialog, QProgressBar,
    QComboBox, QFrame, QDialog, QTableWidget, QTableWidgetItem,
    QHeaderView, QAbstractItemView, QMessageBox, QMenu, QAction,
    QInputDialog, QCheckBox
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal, QSettings
from PyQt5.QtGui import QFont, QDragEnterEvent, QDropEvent, QKeySequence

# ======================== 依赖安全预检 ========================
HAS_DOCX = False
try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    pass

HAS_PPTX = False
try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    pass

HAS_OPENPYXL = False
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    pass

# ======================== 日志配置 ========================
try:
    LOG_DIR = Path("logs")
    LOG_DIR.mkdir(exist_ok=True)
    LOG_FILE = LOG_DIR / "app.log"
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        pass
except (PermissionError, OSError):
    LOG_DIR = Path(os.path.expanduser("~")) / ".qautodoc"
    LOG_DIR.mkdir(exist_ok=True)
    LOG_FILE = LOG_DIR / "app.log"

logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    encoding="utf-8"
)
logger = logging.getLogger(__name__)


# ======================== 核心业务：字段提取器 ========================
class FieldExtractor:
    def __init__(self):
        self.field_pattern = re.compile(r'[\{｛]{2}(.*?)[\}｝]{2}')

    def extract_fields(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.docx' and not HAS_DOCX:
            raise ImportError("缺少依赖: python-docx。请安装: pip install python-docx")
        if ext == '.pptx' and not HAS_PPTX:
            raise ImportError("缺少依赖: python-pptx。请安装: pip install python-pptx")
        if ext == '.xlsx' and not HAS_OPENPYXL:
            raise ImportError("缺少依赖: openpyxl。请安装: pip install openpyxl")

        try:
            if ext == '.docx':
                return self._extract_docx(file_path)
            elif ext == '.pptx':
                return self._extract_pptx(file_path)
            elif ext == '.xlsx':
                return self._extract_xlsx(file_path)
            else:
                raise ValueError(f"不支持的模板格式: {ext}")
        except PermissionError:
            raise Exception("文档正被其他程序(如Office)占用，请关闭后重试")
        except Exception as e:
            err_type = type(e).__name__
            if "BadZip" in err_type or "InvalidFile" in err_type:
                raise Exception("文档损坏或格式异常，请尝试重新另存为")
            logger.exception("提取字段失败")
            raise Exception(f"解析失败: {str(e)}")

    def _extract_docx(self, path):
        doc = Document(path)
        fields, seen = [], set()
        for p in doc.paragraphs:
            self._collect_fields(p.text, fields, seen)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        self._collect_fields(p.text, fields, seen)
        return fields

    def _extract_pptx(self, path):
        prs = Presentation(path)
        fields, seen = [], set()
        for slide in prs.slides:
            self._extract_shapes(slide.shapes, fields, seen)
        return fields

    def _extract_shapes(self, shapes, fields, seen):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                self._extract_shapes(shape.shapes, fields, seen)
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    self._collect_fields(p.text, fields, seen)
            if hasattr(shape, "has_table") and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            self._collect_fields(p.text, fields, seen)

    def _extract_xlsx(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        fields, seen = [], set()
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        self._collect_fields(cell, fields, seen)
        return fields

    def _collect_fields(self, text, fields_list, seen_set):
        if not text:
            return
        for m in self.field_pattern.findall(text):
            f = m.strip()
            if f and f not in seen_set:
                fields_list.append(f)
                seen_set.add(f)


# ======================== 核心业务：文档处理器 ========================
class DocumentProcessor:
    def __init__(self):
        self.field_pattern = re.compile(r'[\{｛]{2}(.*?)[\}｝]{2}')

    def process_document(self, template_path, output_dir, field_values, naming_field=None, index=None):
        if not os.path.exists(template_path):
            raise FileNotFoundError("模板文件不存在")
        if not field_values:
            raise ValueError("当前条目数据为空")

        ext = os.path.splitext(template_path)[1].lower()
        filename = self._gen_filename(template_path, field_values, naming_field, index)
        out_path = os.path.join(output_dir, filename)

        try:
            if ext == '.docx':
                self._process_docx(template_path, out_path, field_values)
            elif ext == '.pptx':
                self._process_pptx(template_path, out_path, field_values)
            elif ext == '.xlsx':
                self._process_xlsx(template_path, out_path, field_values)
            return out_path
        except PermissionError:
            raise Exception(f"目标文件已被独占锁定: {filename}")
        except Exception as e:
            logger.exception("文档处理失败")
            raise Exception(f"构建失败: {str(e)}")

    def _process_docx(self, tpl, out, fv):
        doc = Document(tpl)
        for p in doc.paragraphs:
            self._replace_in_para(p, fv)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        self._replace_in_para(p, fv)
        doc.save(out)

    def _process_pptx(self, tpl, out, fv):
        prs = Presentation(tpl)
        for slide in prs.slides:
            self._process_shapes(slide.shapes, fv)
        prs.save(out)

    def _process_shapes(self, shapes, fv):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                self._process_shapes(shape.shapes, fv)
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    self._replace_in_para(p, fv)
            if hasattr(shape, "has_table") and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            self._replace_in_para(p, fv)

    def _process_xlsx(self, tpl, out, fv):
        wb = openpyxl.load_workbook(tpl)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        matches = list(self.field_pattern.finditer(cell.value))
                        if matches:
                            new_text = cell.value
                            for m in sorted(matches, key=lambda x: x.start(), reverse=True):
                                name = m.group(1).strip()
                                if name in fv:
                                    new_text = new_text[:m.start()] + str(fv[name]) + new_text[m.end():]
                            cell.value = new_text
        wb.save(out)

    def _replace_in_para(self, paragraph, fv):
        if not paragraph.text:
            return
        matches = list(self.field_pattern.finditer(paragraph.text))
        if not matches:
            return
        matches.sort(key=lambda x: x.start(), reverse=True)
        for m in matches:
            name = m.group(1).strip()
            if name in fv:
                self._replace_fmt_runs(paragraph, m.start(), m.end(), str(fv[name]))

    def _replace_fmt_runs(self, paragraph, sp, ep, repl):
        runs = paragraph.runs
        if not runs:
            return
        pos = 0
        affected = []
        for i, r in enumerate(runs):
            run_text = r.text if r.text else ""
            rl = len(run_text)
            if pos < ep and pos + rl > sp:
                rs = max(0, sp - pos)
                re_ = min(rl, ep - pos)
                affected.append((i, rs, re_))
            pos += rl

        if not affected:
            return

        first_idx = affected[0][0]
        affected.sort(key=lambda x: x[0], reverse=True)

        for idx, rs, re_ in affected:
            run = runs[idx]
            run_text = run.text if run.text else ""
            before = run_text[:rs]
            after = run_text[re_:]
            if idx == first_idx:
                run.text = before + repl + after
            else:
                run.text = before + after

    def _gen_filename(self, tpl, fv, nf, idx):
        tn, ext = os.path.splitext(os.path.basename(tpl))
        if nf and nf in fv and fv[nf]:
            base = f"{tn}_{fv[nf]}"
        else:
            base = tn
            for f in ['姓名', '名称', '标题', 'subject', 'title', 'name', 'ID']:
                if f in fv and fv[f]:
                    base = f"{tn}_{fv[f]}"
                    break
            if base == tn:
                base = f"{tn}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

        if idx is not None:
            base += f"_{idx + 1:03d}"
        return self._sanitize(f"{base}{ext}")

    def _sanitize(self, fn):
        for c in ['<', '>', ':', '"', '/', '\\', '|', '?', '*']:
            fn = fn.replace(c, '_')
        fn = "".join(ch for ch in fn if ord(ch) >= 32)
        name, ext = os.path.splitext(fn)
        name = name.strip().rstrip('.')
        return f"{name or 'output'}{ext}"


# ======================== 批量处理线程 ========================
class BatchProcessingThread(QThread):
    progress_updated = pyqtSignal(int, int)
    processing_finished = pyqtSignal(list, list)

    def __init__(self, template_path, output_dir, field_values_list, naming_field=None):
        super().__init__()
        self.template_path = template_path
        self.output_dir = output_dir
        self.field_values_list = field_values_list
        self.naming_field = naming_field
        self._is_cancelled = False

    def cancel(self):
        self._is_cancelled = True

    def run(self):
        processor = DocumentProcessor()
        result_files = []
        failed_items = []
        total = len(self.field_values_list)

        for i, fv in enumerate(self.field_values_list):
            if self._is_cancelled:
                break
            try:
                out = processor.process_document(self.template_path, self.output_dir, fv, self.naming_field, i)
                result_files.append(out)
            except Exception as e:
                failed_items.append(f"第 {i+1} 项: {str(e)}")
            self.progress_updated.emit(i + 1, total)

        self.processing_finished.emit(result_files, failed_items)


# ======================== 查找与替换对话框 ========================
class FindReplaceDialog(QDialog):
    def __init__(self, table_widget, parent=None):
        super().__init__(parent)
        self.table = table_widget
        self.setWindowTitle("查找和替换")
        self.setFixedSize(360, 160)
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowContextHelpButtonHint)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 14, 16, 14)
        layout.setSpacing(10)

        # 查找内容
        r1 = QHBoxLayout()
        r1.addWidget(QLabel("查找内容:"))
        self.find_input = QLineEdit()
        r1.addWidget(self.find_input)
        layout.addLayout(r1)

        # 替换为
        r2 = QHBoxLayout()
        r2.addWidget(QLabel("替换为:  "))
        self.replace_input = QLineEdit()
        r2.addWidget(self.replace_input)
        layout.addLayout(r2)

        # 选项
        self.match_case = QCheckBox("区分大小写")
        layout.addWidget(self.match_case)

        # 按钮条
        btn_box = QHBoxLayout()
        btn_box.addStretch()

        self.btn_replace_all = QPushButton("全部替换")
        self.btn_replace_all.setObjectName("PrimaryBtn")
        self.btn_replace_all.clicked.connect(self.do_replace_all)
        btn_box.addWidget(self.btn_replace_all)

        self.btn_close = QPushButton("关闭")
        self.btn_close.setObjectName("SecondaryBtn")
        self.btn_close.clicked.connect(self.close)
        btn_box.addWidget(self.btn_close)

        layout.addLayout(btn_box)

    def do_replace_all(self):
        find_txt = self.find_input.text()
        repl_txt = self.replace_input.text()
        if not find_txt:
            return

        self.table.record_undo_snapshot()
        count = 0
        cs = Qt.CaseSensitive if self.match_case.isChecked() else Qt.CaseInsensitive

        for r in range(self.table.rowCount()):
            for c in range(self.table.columnCount()):
                item = self.table.item(r, c)
                if item and item.text():
                    text = item.text()
                    if cs == Qt.CaseInsensitive:
                        pattern = re.compile(re.escape(find_txt), re.IGNORECASE)
                        new_text, n = pattern.subn(repl_txt, text)
                    else:
                        n = text.count(find_txt)
                        new_text = text.replace(find_txt, repl_txt)
                    if n > 0:
                        item.setText(new_text)
                        count += n

        QMessageBox.information(self, "替换完成", f"已完成匹配并成功替换了 {count} 处内容！")


# ======================== 核心重构：Excel级全功能数据网格 ========================
class ExcelDataGrid(QWidget):
    dataImported = pyqtSignal(int)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.fields = []
        self.undo_stack = []
        self.redo_stack = []
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(5)

        # 顶部 Excel 工具条
        toolbar = QHBoxLayout()
        toolbar.setSpacing(6)

        title_lbl = QLabel("<b>数据矩阵</b>")
        title_lbl.setStyleSheet("font-size: 13px; color: #0f172a;")
        toolbar.addWidget(title_lbl)

        # 常用快速操作按钮
        self.btn_add_row = QPushButton("➕ 行")
        self.btn_add_row.setToolTip("在底部追加一行")
        self.btn_add_row.setObjectName("SecondaryBtn")
        self.btn_add_row.clicked.connect(self.add_empty_row)
        toolbar.addWidget(self.btn_add_row)

        self.btn_del_row = QPushButton("➖ 删行")
        self.btn_del_row.setToolTip("删除选中行")
        self.btn_del_row.setObjectName("SecondaryBtn")
        self.btn_del_row.clicked.connect(self.remove_selected_rows)
        toolbar.addWidget(self.btn_del_row)

        self.btn_fill_down = QPushButton("⬇️ 向下填充 (Ctrl+D)")
        self.btn_fill_down.setToolTip("将选中区域首行复制填充至其余选中单元格")
        self.btn_fill_down.setObjectName("SecondaryBtn")
        self.btn_fill_down.clicked.connect(self.fill_down)
        toolbar.addWidget(self.btn_fill_down)

        self.btn_seq = QPushButton("🔢 序列自增")
        self.btn_seq.setToolTip("将选中的首个单元格数字以 1, 2, 3... 规律向下递增填充")
        self.btn_seq.setObjectName("SecondaryBtn")
        self.btn_seq.clicked.connect(self.fill_sequence)
        toolbar.addWidget(self.btn_seq)

        self.btn_find = QPushButton("🔍 查找替换")
        self.btn_find.setObjectName("SecondaryBtn")
        self.btn_find.clicked.connect(self.open_find_replace)
        toolbar.addWidget(self.btn_find)

        toolbar.addStretch()

        # 数据清洗下拉功能
        self.btn_clean_empty = QPushButton("🧹 剔除空行")
        self.btn_clean_empty.setObjectName("SecondaryBtn")
        self.btn_clean_empty.clicked.connect(self.clean_empty_rows)
        toolbar.addWidget(self.btn_clean_empty)

        self.btn_dedup = QPushButton("✂️ 去除重复")
        self.btn_dedup.setObjectName("SecondaryBtn")
        self.btn_dedup.clicked.connect(self.remove_duplicates)
        toolbar.addWidget(self.btn_dedup)

        self.btn_import = QPushButton("📂 导入 Excel/CSV")
        self.btn_import.setObjectName("SecondaryBtn")
        self.btn_import.clicked.connect(self.browse_external_file)
        toolbar.addWidget(self.btn_import)

        layout.addLayout(toolbar)

        # 核心表格
        self.table = QTableWidget(0, 0)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self.show_context_menu)
        self.table.itemSelectionChanged.connect(self.update_selection_metrics)

        # 表头配置：支持自由拖动拉伸与点击排序
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setMinimumSectionSize(120)
        header.setFixedHeight(34)
        header.setStretchLastSection(True)
        header.sectionClicked.connect(self.sort_by_column)

        self.table.verticalHeader().setDefaultSectionSize(28)
        self.table.verticalHeader().setMinimumSectionSize(24)
        self.table.setStyleSheet("""
            QTableWidget {
                background: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 6px;
                gridline-color: #f1f5f9;
                font-size: 12px;
            }
            QTableWidget::item { padding: 3px 6px; }
            QTableWidget::item:selected { background: #dbeafe; color: #1e3a8a; }
            QHeaderView::section {
                background: #f8fafc;
                color: #1e293b;
                font-weight: bold;
                border: none;
                border-right: 1px solid #e2e8f0;
                border-bottom: 2px solid #cbd5e1;
                padding: 4px 8px;
                font-size: 12px;
            }
        """)
        layout.addWidget(self.table, 1)

        # Excel 经典底部微型统计栏
        self.stat_bar = QHBoxLayout()
        self.stat_bar.setContentsMargins(4, 2, 4, 2)
        self.stat_info = QLabel("共 0 行 0 列")
        self.stat_info.setStyleSheet("color: #64748b; font-size: 11px;")
        self.stat_metrics = QLabel("")
        self.stat_metrics.setStyleSheet("color: #0284c7; font-size: 11px; font-weight: 500;")

        self.stat_bar.addWidget(self.stat_info)
        self.stat_bar.addStretch()
        self.stat_bar.addWidget(self.stat_metrics)
        layout.addLayout(self.stat_bar)

    # ------------------ 撤销与重做快照栈 (Undo / Redo) ------------------
    def record_undo_snapshot(self):
        snapshot = []
        for r in range(self.table.rowCount()):
            row_data = [self.table.item(r, c).text() if self.table.item(r, c) else "" 
                        for c in range(self.table.columnCount())]
            snapshot.append(row_data)
        self.undo_stack.append(snapshot)
        if len(self.undo_stack) > 30:
            self.undo_stack.pop(0)
        self.redo_stack.clear()

    def undo(self):
        if not self.undo_stack:
            return
        # 保存当前作为 redo
        current_data = []
        for r in range(self.table.rowCount()):
            row_data = [self.table.item(r, c).text() if self.table.item(r, c) else "" 
                        for c in range(self.table.columnCount())]
            current_data.append(row_data)
        self.redo_stack.append(current_data)

        last_state = self.undo_stack.pop()
        self._apply_snapshot(last_state)

    def redo(self):
        if not self.redo_stack:
            return
        # 保存当前作为 undo
        current_data = []
        for r in range(self.table.rowCount()):
            row_data = [self.table.item(r, c).text() if self.table.item(r, c) else "" 
                        for c in range(self.table.columnCount())]
            current_data.append(row_data)
        self.undo_stack.append(current_data)

        next_state = self.redo_stack.pop()
        self._apply_snapshot(next_state)

    def _apply_snapshot(self, data):
        self.table.setRowCount(len(data))
        for r, row_values in enumerate(data):
            for c, val in enumerate(row_values):
                if c < self.table.columnCount():
                    self.table.setItem(r, c, QTableWidgetItem(val))
        self.update_row_count_display()

    # ------------------ Excel 经典快捷键截获 ------------------
    def keyPressEvent(self, event):
        # 撤销 Ctrl+Z
        if event.matches(QKeySequence.Undo):
            self.undo()
            return
        # 重做 Ctrl+Y
        if event.matches(QKeySequence.Redo):
            self.redo()
            return
        # 复制 Ctrl+C
        if event.matches(QKeySequence.Copy):
            self.copy_selection()
            return
        # 剪切 Ctrl+X
        if event.matches(QKeySequence.Cut):
            self.cut_selection()
            return
        # 粘贴 Ctrl+V
        if event.matches(QKeySequence.Paste):
            self.paste_from_clipboard()
            return
        # 向下填充 Ctrl+D
        if event.modifiers() == Qt.ControlModifier and event.key() == Qt.Key_D:
            self.fill_down()
            return
        # 查找替换 Ctrl+F 或 Ctrl+H
        if event.modifiers() == Qt.ControlModifier and event.key() in (Qt.Key_F, Qt.Key_H):
            self.open_find_replace()
            return
        # 删除单元格内容 Delete / Backspace
        if event.key() in (Qt.Key_Delete, Qt.Key_Backspace):
            self.delete_selected_cells()
            return

        super().keyPressEvent(event)

    # ------------------ 剪贴板矩阵读写 ------------------
    def copy_selection(self):
        ranges = self.table.selectedRanges()
        if not ranges:
            return
        selected_range = ranges[0]
        copied_text = ""
        for r in range(selected_range.topRow(), selected_range.bottomRow() + 1):
            row_items = []
            for c in range(selected_range.leftColumn(), selected_range.rightColumn() + 1):
                item = self.table.item(r, c)
                row_items.append(item.text() if item else "")
            copied_text += "\t".join(row_items) + "\n"
        QApplication.clipboard().setText(copied_text.rstrip("\n"))

    def cut_selection(self):
        self.record_undo_snapshot()
        self.copy_selection()
        self.delete_selected_cells()

    def paste_from_clipboard(self):
        text = QApplication.clipboard().text()
        if not text:
            return

        self.record_undo_snapshot()
        lines = text.split("\n")
        while lines and not lines[-1].strip():
            lines.pop()
        if not lines:
            return

        indexes = self.table.selectedIndexes()
        start_row = min([i.row() for i in indexes]) if indexes else 0
        start_col = min([i.column() for i in indexes]) if indexes else 0

        for r_offset, line in enumerate(lines):
            tr = start_row + r_offset
            if tr >= self.table.rowCount():
                self.table.insertRow(self.table.rowCount())
            cols = line.split("\t")
            for c_offset, val in enumerate(cols):
                tc = start_col + c_offset
                if tc < self.table.columnCount():
                    self.table.setItem(tr, tc, QTableWidgetItem(val.strip("\r")))

        self.update_row_count_display()

    def delete_selected_cells(self):
        self.record_undo_snapshot()
        for item in self.table.selectedItems():
            item.setText("")

    # ------------------ Excel 级特性：填充、序列与清洗 ------------------
    def fill_down(self):
        """向下填充 (Ctrl+D)"""
        ranges = self.table.selectedRanges()
        if not ranges:
            return
        self.record_undo_snapshot()
        for sel in ranges:
            if sel.rowCount() <= 1:
                continue
            for c in range(sel.leftColumn(), sel.rightColumn() + 1):
                top_item = self.table.item(sel.topRow(), c)
                val = top_item.text() if top_item else ""
                for r in range(sel.topRow() + 1, sel.bottomRow() + 1):
                    self.table.setItem(r, c, QTableWidgetItem(val))

    def fill_sequence(self):
        """智能序列自增"""
        ranges = self.table.selectedRanges()
        if not ranges:
            return
        self.record_undo_snapshot()
        for sel in ranges:
            if sel.rowCount() <= 1:
                continue
            for c in range(sel.leftColumn(), sel.rightColumn() + 1):
                top_item = self.table.item(sel.topRow(), c)
                base_txt = top_item.text() if top_item else "1"
                
                # 正则分离前缀与末尾数字
                m = re.search(r'^(.*?)(\d+)$', base_txt)
                if m:
                    prefix, num_str = m.group(1), m.group(2)
                    num_val = int(num_str)
                    pad_len = len(num_str)
                    for idx, r in enumerate(range(sel.topRow() + 1, sel.bottomRow() + 1), start=1):
                        next_num = f"{num_val + idx:0{pad_len}d}"
                        self.table.setItem(r, c, QTableWidgetItem(f"{prefix}{next_num}"))
                else:
                    # 如果不是数字结尾，退化为纯复制
                    for r in range(sel.topRow() + 1, sel.bottomRow() + 1):
                        self.table.setItem(r, c, QTableWidgetItem(base_txt))

    def fill_column_constant(self, col):
        """整列批量设置相同值"""
        col_name = self.fields[col] if col < len(self.fields) else f"第{col+1}列"
        val, ok = QInputDialog.getText(self, "整列覆写", f"请输入要批量填入「{col_name}」的内容:")
        if ok and val is not None:
            self.record_undo_snapshot()
            for r in range(self.table.rowCount()):
                self.table.setItem(r, col, QTableWidgetItem(val))

    def clean_empty_rows(self):
        """剔除所有全空行"""
        self.record_undo_snapshot()
        removed = 0
        for r in range(self.table.rowCount() - 1, -1, -1):
            is_empty = True
            for c in range(self.table.columnCount()):
                item = self.table.item(r, c)
                if item and item.text().strip():
                    is_empty = False
                    break
            if is_empty:
                self.table.removeRow(r)
                removed += 1
        self.update_row_count_display()
        QMessageBox.information(self, "清洗完成", f"已成功移除 {removed} 个全空行！")

    def remove_duplicates(self):
        """整行去重"""
        self.record_undo_snapshot()
        seen = set()
        removed = 0
        for r in range(self.table.rowCount() - 1, -1, -1):
            row_vals = tuple(self.table.item(r, c).text().strip() if self.table.item(r, c) else "" 
                             for c in range(self.table.columnCount()))
            if row_vals in seen:
                self.table.removeRow(r)
                removed += 1
            else:
                seen.add(row_vals)
        self.update_row_count_display()
        QMessageBox.information(self, "去重完成", f"已扫描并移除了 {removed} 条完全重复的数据！")

    def sort_by_column(self, col):
        """点击表头升序/降序"""
        order = getattr(self, f"_sort_order_{col}", Qt.AscendingOrder)
        self.table.sortItems(col, order)
        setattr(self, f"_sort_order_{col}", Qt.DescendingOrder if order == Qt.AscendingOrder else Qt.AscendingOrder)

    def open_find_replace(self):
        FindReplaceDialog(self, self).exec_()

    # ------------------ Excel 右键上下文菜单 ------------------
    def show_context_menu(self, pos):
        menu = QMenu(self)
        
        act_copy = menu.addAction("📋 复制 (Ctrl+C)")
        act_copy.triggered.connect(self.copy_selection)

        act_cut = menu.addAction("✂️ 剪切 (Ctrl+X)")
        act_cut.triggered.connect(self.cut_selection)

        act_paste = menu.addAction("📥 粘贴 (Ctrl+V)")
        act_paste.triggered.connect(self.paste_from_clipboard)

        act_del = menu.addAction("🗑️ 清空单元格 (Delete)")
        act_del.triggered.connect(self.delete_selected_cells)

        menu.addSeparator()

        act_fill_down = menu.addAction("⬇️ 向下填充 (Ctrl+D)")
        act_fill_down.triggered.connect(self.fill_down)

        act_seq = menu.addAction("🔢 序列自增")
        act_seq.triggered.connect(self.fill_sequence)

        item = self.table.itemAt(pos)
        if item:
            col = item.column()
            act_col_const = menu.addAction(f"🏷️ 整列统一填入相同值...")
            act_col_const.triggered.connect(lambda: self.fill_column_constant(col))

        menu.addSeparator()

        act_ins_above = menu.addAction("➕ 在上方插入行")
        act_ins_above.triggered.connect(self.insert_row_above)

        act_ins_below = menu.addAction("➕ 在下方插入行")
        act_ins_below.triggered.connect(self.insert_row_below)

        act_del_row = menu.addAction("➖ 删除选中整行")
        act_del_row.triggered.connect(self.remove_selected_rows)

        menu.addSeparator()
        act_fr = menu.addAction("🔍 查找和替换 (Ctrl+F)")
        act_fr.triggered.connect(self.open_find_replace)

        menu.exec_(self.table.mapToGlobal(pos))

    def insert_row_above(self):
        self.record_undo_snapshot()
        idx = self.table.currentRow()
        self.table.insertRow(max(0, idx))
        self.update_row_count_display()

    def insert_row_below(self):
        self.record_undo_snapshot()
        idx = self.table.currentRow()
        self.table.insertRow(idx + 1 if idx >= 0 else self.table.rowCount())
        self.update_row_count_display()

    # ------------------ Excel 经典底栏数据统计 ------------------
    def update_selection_metrics(self):
        selected_items = self.table.selectedItems()
        count = len(selected_items)
        if count == 0:
            self.stat_metrics.setText("")
            return

        non_empty = 0
        numeric_vals = []

        for it in selected_items:
            t = it.text().strip()
            if t:
                non_empty += 1
                try:
                    num = float(t)
                    numeric_vals.append(num)
                except ValueError:
                    pass

        msg = f"选中: {count} 个单元格 | 非空: {non_empty}"
        if len(numeric_vals) >= 2:
            s = sum(numeric_vals)
            avg = s / len(numeric_vals)
            msg += f" | 求和: {s:g} | 平均: {avg:.2f}"
        self.stat_metrics.setText(msg)

    def update_row_count_display(self):
        self.stat_info.setText(f"共 {self.table.rowCount()} 行 {self.table.columnCount()} 列")

    # ------------------ 外部 Excel/CSV 拖放与解析 ------------------
    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            path = event.mimeData().urls()[0].toLocalFile().lower()
            if path.endswith(('.xlsx', '.xls', '.csv')):
                event.acceptProposedAction()
                return
        event.ignore()

    def dropEvent(self, event: QDropEvent):
        urls = event.mimeData().urls()
        if urls:
            self.load_external_file(urls[0].toLocalFile())

    def browse_external_file(self):
        fp, _ = QFileDialog.getOpenFileName(self, "选择表格文件", "", "数据表格 (*.xlsx *.xls *.csv)")
        if fp:
            self.load_external_file(fp)

    def load_external_file(self, file_path):
        ext = os.path.splitext(file_path)[1].lower()
        records = []
        try:
            if ext == '.xlsx':
                if not HAS_OPENPYXL:
                    QMessageBox.warning(self, "依赖缺失", "未检测到 openpyxl，无法解析 Excel！")
                    return
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet = wb.active
                rows = list(sheet.iter_rows(values_only=True))
                if len(rows) >= 2:
                    headers = [str(c).strip() if c is not None else "" for c in rows[0]]
                    for row in rows[1:]:
                        if any(cell is not None and str(cell).strip() != "" for cell in row):
                            rec = {headers[i]: str(row[i]).strip() if i < len(row) and row[i] is not None else "" 
                                   for i, h in enumerate(headers) if h}
                            records.append(rec)

            elif ext == '.csv':
                for enc in ['utf-8-sig', 'gb18030', 'utf-8', 'ansi']:
                    try:
                        with open(file_path, 'r', encoding=enc) as f:
                            records = [{k.strip(): v.strip() if v else "" for k, v in r.items() if k} 
                                       for r in csv.DictReader(f)]
                        break
                    except UnicodeDecodeError:
                        continue
        except Exception as e:
            QMessageBox.critical(self, "读取失败", f"无法解析文件: {str(e)}")
            return

        if not records:
            QMessageBox.warning(self, "无数据", "未能成功提取出有效的数据行！")
            return

        self.record_undo_snapshot()
        if not self.fields:
            self.set_fields(list(records[0].keys()))

        self.populate_data(records)
        self.dataImported.emit(len(records))
        QMessageBox.information(self, "导入成功", f"成功将 {len(records)} 行数据灌入矩阵！")

    # ------------------ 基础增删与数据导出 ------------------
    def set_fields(self, fields):
        self.fields = fields
        self.table.clear()
        self.table.setColumnCount(len(fields))
        self.table.setHorizontalHeaderLabels(fields)
        header = self.table.horizontalHeader()
        for i, f in enumerate(fields):
            header.resizeSection(i, max(130, len(f) * 16 + 32))
        if fields and self.table.rowCount() == 0:
            self.add_empty_row()
        self.update_row_count_display()

    def add_empty_row(self):
        if not self.fields:
            return
        self.record_undo_snapshot()
        row = self.table.rowCount()
        self.table.insertRow(row)
        for col in range(len(self.fields)):
            self.table.setItem(row, col, QTableWidgetItem(""))
        self.update_row_count_display()

    def remove_selected_rows(self):
        selected_indexes = self.table.selectionModel().selectedRows()
        if not selected_indexes:
            if self.table.rowCount() > 0:
                self.record_undo_snapshot()
                self.table.removeRow(self.table.rowCount() - 1)
                self.update_row_count_display()
            return
        self.record_undo_snapshot()
        for index in sorted(selected_indexes, reverse=True):
            self.table.removeRow(index.row())
        self.update_row_count_display()

    def get_data(self):
        data = []
        rows = self.table.rowCount()
        cols = self.table.columnCount()
        for r in range(rows):
            row_dict = {}
            has_val = False
            for c in range(cols):
                item = self.table.item(r, c)
                val = item.text().strip() if item else ""
                row_dict[self.fields[c]] = val
                if val:
                    has_val = True
            if has_val:
                data.append(row_dict)
        return data

    def populate_data(self, records):
        if not self.fields or not records:
            return
        self.table.setRowCount(0)
        for r, rec in enumerate(records):
            self.table.insertRow(r)
            for c, field in enumerate(self.fields):
                val = rec.get(field, "")
                if not val:
                    for k, v in rec.items():
                        if k.replace('{', '').replace('}', '').strip() == field:
                            val = v
                            break
                self.table.setItem(r, c, QTableWidgetItem(str(val)))
        self.update_row_count_display()


# ======================== 模板卡片 ========================
class CompactTemplateBar(QFrame):
    fileDropped = pyqtSignal(str)
    fileCleared = pyqtSignal()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.setObjectName("TemplateBar")
        self.loaded_path = ""
        self.init_ui()

    def init_ui(self):
        self.main_layout = QHBoxLayout(self)
        self.main_layout.setContentsMargins(12, 6, 12, 6)
        self.main_layout.setSpacing(10)

        self.empty_widget = QWidget()
        ew_lay = QHBoxLayout(self.empty_widget)
        ew_lay.setContentsMargins(0, 0, 0, 0)
        ew_lay.setSpacing(8)

        self.icon_label = QLabel("📥")
        self.icon_label.setStyleSheet("font-size: 16px; background: transparent;")
        ew_lay.addWidget(self.icon_label)

        self.tip_label = QLabel("<b>步骤 1：</b>拖入 DOCX / PPTX / XLSX 模板文档，或点击浏览")
        self.tip_label.setStyleSheet("font-size: 12px; color: #334155;")
        ew_lay.addWidget(self.tip_label, 1)

        self.select_btn = QPushButton("浏览模板")
        self.select_btn.setObjectName("SecondaryBtn")
        self.select_btn.setCursor(Qt.PointingHandCursor)
        self.select_btn.clicked.connect(self._select_file_dialog)
        ew_lay.addWidget(self.select_btn)

        self.main_layout.addWidget(self.empty_widget)

        self.loaded_widget = QWidget()
        self.loaded_widget.setVisible(False)
        lw_lay = QHBoxLayout(self.loaded_widget)
        lw_lay.setContentsMargins(0, 0, 0, 0)
        lw_lay.setSpacing(8)

        self.badge = QLabel("DOCX")
        self.badge.setStyleSheet("background: #2563eb; color: white; border-radius: 4px; padding: 2px 6px; font-size: 11px; font-weight: bold;")
        lw_lay.addWidget(self.badge)

        self.file_name_lbl = QLabel("template.docx")
        self.file_name_lbl.setStyleSheet("font-weight: bold; font-size: 12px; color: #0f172a;")
        lw_lay.addWidget(self.file_name_lbl)

        self.fields_summary_lbl = QLabel("")
        self.fields_summary_lbl.setStyleSheet("font-size: 12px; color: #64748b;")
        lw_lay.addWidget(self.fields_summary_lbl, 1)

        self.change_btn = QPushButton("更换")
        self.change_btn.setObjectName("SecondaryBtn")
        self.change_btn.setCursor(Qt.PointingHandCursor)
        self.change_btn.clicked.connect(self._select_file_dialog)
        lw_lay.addWidget(self.change_btn)

        self.del_btn = QPushButton("清除")
        self.del_btn.setObjectName("DangerBtn")
        self.del_btn.setCursor(Qt.PointingHandCursor)
        self.del_btn.clicked.connect(self.clear_template)
        lw_lay.addWidget(self.del_btn)

        self.main_layout.addWidget(self.loaded_widget)
        self._set_style(False)

    def _set_style(self, active):
        bc = "#2563eb" if active else "#cbd5e1"
        bg = "#eff6ff" if active else "#f8fafc"
        self.setStyleSheet(f"""
            #TemplateBar {{
                background-color: {bg};
                border: 1px dashed {bc};
                border-radius: 6px;
            }}
            #TemplateBar:hover {{
                border-color: #3b82f6;
                background-color: #f1f5f9;
            }}
        """)

    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self._set_style(True)

    def dragLeaveEvent(self, event):
        self._set_style(False)

    def dropEvent(self, event: QDropEvent):
        self._set_style(False)
        urls = event.mimeData().urls()
        if urls:
            path = urls[0].toLocalFile()
            if path.lower().endswith(('.docx', '.pptx', '.xlsx')):
                self.fileDropped.emit(path)
            else:
                QMessageBox.warning(self, "格式不符", "仅支持 .docx / .pptx / .xlsx 模板！")

    def mousePressEvent(self, event):
        if not self.loaded_path and event.button() == Qt.LeftButton:
            self._select_file_dialog()
        super().mousePressEvent(event)

    def _select_file_dialog(self):
        fp, _ = QFileDialog.getOpenFileName(self, "选择模板文档", "", "模板 (*.docx *.pptx *.xlsx)")
        if fp:
            self.fileDropped.emit(fp)

    def set_file_loaded(self, path, fields):
        self.loaded_path = path
        fn = os.path.basename(path)
        ext = os.path.splitext(fn)[1].lower().replace('.', '').upper()
        self.file_name_lbl.setText(fn)

        colors = {'DOCX': '#0284c7', 'PPTX': '#ea580c', 'XLSX': '#16a34a'}
        self.badge.setText(ext)
        self.badge.setStyleSheet(f"background: {colors.get(ext, '#64748b')}; color: white; border-radius: 4px; padding: 2px 6px; font-size: 11px; font-weight: bold;")

        if fields:
            self.fields_summary_lbl.setText(f"识别出 {len(fields)} 个占位变量")
        else:
            self.fields_summary_lbl.setText("⚠️ 未检测到 {{变量}} 占位符")

        self.empty_widget.setVisible(False)
        self.loaded_widget.setVisible(True)

    def clear_template(self):
        self.loaded_path = ""
        self.empty_widget.setVisible(True)
        self.loaded_widget.setVisible(False)
        self.fileCleared.emit()


# ======================== 主界面 ========================
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.settings = QSettings("QAutoDoc", "ExcelEdition")
        self.template_path = ""
        self.extracted_fields = []
        self.batch_thread = None

        self.init_responsive_geometry()
        self.init_theme()
        self.init_ui()

    def init_responsive_geometry(self):
        screen = QApplication.primaryScreen()
        if screen:
            avail = screen.availableGeometry()
            w = max(800, min(int(avail.width() * 0.62), 1140))
            h = max(560, min(int(avail.height() * 0.70), 780))
            x = avail.x() + (avail.width() - w) // 2
            y = avail.y() + (avail.height() - h) // 2
            self.setGeometry(x, y, w, h)
        else:
            self.resize(920, 640)

        self.setMinimumSize(760, 500)
        self.setWindowTitle("QAutoDoc 智能文档批量生成器")

    def init_theme(self):
        self.setStyleSheet("""
            QWidget {
                font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "Microsoft YaHei", sans-serif;
                color: #0f172a;
            }
            QMainWindow { background-color: #ffffff; }
            QLineEdit, QComboBox {
                padding: 4px 8px;
                border: 1px solid #cbd5e1;
                border-radius: 4px;
                background: #ffffff;
                font-size: 12px;
            }
            QComboBox::drop-down { border: none; width: 16px; }
            QPushButton#PrimaryBtn {
                background-color: #2563eb;
                color: #ffffff;
                font-size: 12px;
                font-weight: 600;
                border: none;
                border-radius: 4px;
                padding: 5px 14px;
            }
            QPushButton#PrimaryBtn:hover { background-color: #1d4ed8; }
            QPushButton#PrimaryBtn:disabled { background-color: #94a3b8; }
            QPushButton#SecondaryBtn {
                background-color: #ffffff;
                color: #334155;
                font-size: 12px;
                border: 1px solid #cbd5e1;
                border-radius: 4px;
                padding: 4px 8px;
            }
            QPushButton#SecondaryBtn:hover { background-color: #f1f5f9; }
            QPushButton#DangerBtn {
                background-color: #fee2e2;
                color: #dc2626;
                font-size: 12px;
                border: 1px solid #fecaca;
                border-radius: 4px;
                padding: 4px 8px;
            }
            QPushButton#DangerBtn:hover { background-color: #fecaca; }
            QProgressBar {
                border: none;
                background-color: #e2e8f0;
                border-radius: 3px;
                height: 5px;
            }
            QProgressBar::chunk { background-color: #2563eb; border-radius: 3px; }
            QMenu {
                background-color: #ffffff;
                border: 1px solid #cbd5e1;
                border-radius: 4px;
                padding: 4px 0;
            }
            QMenu::item { padding: 5px 20px; font-size: 12px; color: #334155; }
            QMenu::item:selected { background-color: #e0f2fe; color: #0369a1; }
        """)

    def init_ui(self):
        root = QWidget()
        layout = QVBoxLayout(root)
        layout.setContentsMargins(14, 10, 14, 10)
        layout.setSpacing(8)
        self.setCentralWidget(root)

        # 顶栏
        top_bar = QHBoxLayout()
        top_bar.addWidget(QLabel("📄 <b>QAutoDoc</b> <span style='color:#64748b; font-size:12px;'>批量文档生成工作站</span>"))
        top_bar.addStretch()

        self.export_tpl_btn = QPushButton("⬇️ 下载 Excel 表头模板")
        self.export_tpl_btn.setObjectName("SecondaryBtn")
        self.export_tpl_btn.clicked.connect(self.export_excel_template)
        top_bar.addWidget(self.export_tpl_btn)

        layout.addLayout(top_bar)

        # 模板区域
        self.template_bar = CompactTemplateBar()
        self.template_bar.fileDropped.connect(self.on_template_selected)
        self.template_bar.fileCleared.connect(self.on_template_cleared)
        layout.addWidget(self.template_bar)

        # 核心：Excel 级数据矩阵
        self.grid_view = ExcelDataGrid()
        self.grid_view.dataImported.connect(lambda n: self.show_status(f"成功导入 {n} 行数据", "#2563eb"))
        layout.addWidget(self.grid_view, 1)

        # 底栏
        footer = QFrame()
        footer.setStyleSheet("border-top: 1px solid #e2e8f0; padding-top: 4px;")
        f_layout = QHBoxLayout(footer)
        f_layout.setContentsMargins(0, 4, 0, 0)
        f_layout.setSpacing(10)

        f_layout.addWidget(QLabel("命名依据:"))
        self.naming_combo = QComboBox()
        self.naming_combo.addItem("[智能默认]")
        self.naming_combo.setMinimumWidth(120)
        f_layout.addWidget(self.naming_combo)

        self.status_msg = QLabel("等待载入模板...")
        self.status_msg.setStyleSheet("color: #64748b; font-size: 12px;")
        f_layout.addWidget(self.status_msg, 1)

        self.progress_bar = QProgressBar()
        self.progress_bar.setFixedWidth(100)
        self.progress_bar.setVisible(False)
        f_layout.addWidget(self.progress_bar)

        self.open_dir_btn = QPushButton("📂 打开输出文件夹")
        self.open_dir_btn.setObjectName("SecondaryBtn")
        self.open_dir_btn.setVisible(False)
        self.open_dir_btn.clicked.connect(self.open_output_folder)
        f_layout.addWidget(self.open_dir_btn)

        self.action_btn = QPushButton("🚀 开始批量生成")
        self.action_btn.setObjectName("PrimaryBtn")
        self.action_btn.setFixedHeight(32)
        self.action_btn.setCursor(Qt.PointingHandCursor)
        self.action_btn.clicked.connect(self.handle_action_btn)
        f_layout.addWidget(self.action_btn)

        layout.addWidget(footer)

    def show_status(self, text, color="#64748b"):
        self.status_msg.setText(text)
        self.status_msg.setStyleSheet(f"color: {color}; font-size: 12px; font-weight: bold;")

    def on_template_selected(self, path):
        try:
            fields = FieldExtractor().extract_fields(path)
            self.template_path = path
            self.extracted_fields = fields
            self.template_bar.set_file_loaded(path, fields)

            self.grid_view.set_fields(fields)

            self.naming_combo.clear()
            self.naming_combo.addItem("[智能默认]")
            for f in fields:
                self.naming_combo.addItem(f"以 {f} 命名")

            self.show_status(f"模板就绪 · 捕获 {len(fields)} 个占位变量", "#16a34a")
        except Exception as e:
            QMessageBox.critical(self, "解析出错", str(e))
            self.template_bar.clear_template()

    def on_template_cleared(self):
        self.template_path = ""
        self.extracted_fields = []
        self.grid_view.set_fields([])
        self.naming_combo.clear()
        self.naming_combo.addItem("[智能默认]")
        self.show_status("模板已清除")

    def export_excel_template(self):
        if not self.extracted_fields:
            QMessageBox.warning(self, "提示", "请先在上方载入模板文件，以便提取列头！")
            return
        save_path, _ = QFileDialog.getSaveFileName(self, "下载数据模板 Excel", "数据导入模板.xlsx", "Excel 工作簿 (*.xlsx)")
        if save_path:
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "数据导入"
                ws.append(self.extracted_fields)
                wb.save(save_path)
                QMessageBox.information(self, "成功", f"Excel 表头模板已保存至:\n{save_path}")
            except Exception as e:
                QMessageBox.critical(self, "导出受阻", str(e))

    def handle_action_btn(self):
        if self.batch_thread and self.batch_thread.isRunning():
            self.batch_thread.cancel()
            self.action_btn.setEnabled(False)
            self.show_status("正在终止生成...", "#dc2626")
            return
        self.start_batch_generation()

    def start_batch_generation(self):
        if not self.template_path:
            QMessageBox.warning(self, "未选模板", "请先在上方拖入 Word/PPT/Excel 模板！")
            return

        dataset = self.grid_view.get_data()
        if not dataset:
            QMessageBox.warning(self, "数据为空", "表格内无有效数据！请录入或拖入 Excel 数据。")
            return

        out_dir = QFileDialog.getExistingDirectory(self, "选择保存目录")
        if not out_dir:
            return
        self.last_output_dir = out_dir

        naming_idx = self.naming_combo.currentIndex()
        naming_field = self.extracted_fields[naming_idx - 1] if naming_idx > 0 else None

        self.action_btn.setText("⏹ 取消")
        self.action_btn.setObjectName("DangerBtn")
        self.action_btn.setStyleSheet("")
        self.progress_bar.setVisible(True)
        self.progress_bar.setRange(0, len(dataset))
        self.progress_bar.setValue(0)
        self.open_dir_btn.setVisible(False)
        self.show_status(f"开始渲染 (共 {len(dataset)} 份)...", "#2563eb")

        self.batch_thread = BatchProcessingThread(self.template_path, out_dir, dataset, naming_field)
        self.batch_thread.progress_updated.connect(self.on_thread_progress)
        self.batch_thread.processing_finished.connect(self.on_thread_finished)
        self.batch_thread.start()

    def on_thread_progress(self, current, total):
        self.progress_bar.setValue(current)
        self.show_status(f"构建进度 [{current}/{total}]...", "#2563eb")

    def on_thread_finished(self, successes, failures):
        self.action_btn.setEnabled(True)
        self.action_btn.setText("🚀 开始批量生成")
        self.action_btn.setObjectName("PrimaryBtn")
        self.action_btn.setStyleSheet("")
        self.progress_bar.setVisible(False)

        if successes:
            self.open_dir_btn.setVisible(True)

        if not failures:
            self.show_status(f"🎉 全部生成成功 (共 {len(successes)} 份)", "#16a34a")
            QMessageBox.information(self, "完成", f"已成功构建 {len(successes)} 份文件！\n目录: {self.last_output_dir}")
        else:
            self.show_status(f"⚠️ 成功 {len(successes)} 份，异常 {len(failures)} 份", "#d97706")
            QMessageBox.warning(self, "部分构建异常", "错误列表:\n" + "\n".join(failures[:6]))

    def open_output_folder(self):
        if hasattr(self, 'last_output_dir') and os.path.exists(self.last_output_dir):
            sys_name = platform.system()
            if sys_name == "Windows":
                os.startfile(self.last_output_dir)
            elif sys_name == "Darwin":
                subprocess.Popen(["open", self.last_output_dir])
            else:
                subprocess.Popen(["xdg-open", self.last_output_dir])


# ======================== 程序入口 ========================
def main():
    if hasattr(Qt, 'AA_EnableHighDpiScaling'):
        QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    if hasattr(Qt, 'AA_UseHighDpiPixmaps'):
        QApplication.setAttribute(Qt.AA_UseHighDpiPixmaps, True)

    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setFont(QFont("Microsoft YaHei", 9))

    window = MainWindow()
    window.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()